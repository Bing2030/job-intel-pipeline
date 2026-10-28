#!/usr/bin/env python3
# Build a technical + management capability demand map from structured JD data (jobs_structured.csv).
import argparse
import re
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import jobintel_paths as jp

CAPS=[
 # id, category, capability, keywords (regex, case-insensitive; matched against responsibilities + requirements + nice-to-have)
 ('agent_runtime','技术','Agent 系统设计/运行时','agent|智能体|multi-agent|多智能体|react|任务规划|planning|工具调用|tool calling|记忆|memory|workflow|编排|长任务|long-horizon|durable|harness|runtime'),
 ('agent_platform','技术','Agent 平台治理','agent平台|agent 平台|注册|发现|权限|租户|审计|配额|发布|版本|skill|技能|marketplace|治理'),
 ('harness_runtime','技术','Harness / Runtime / 上下文','harness|runtime|sandbox|沙箱|上下文|context|tool protocol|mcp|langgraph|dify|coze'),
 ('rag','技术','RAG / 知识检索','rag|检索|向量|embedding|rerank|知识库|混合检索'),
 ('evaluation','技术','Evaluation / 可观测','评测|evaluation|llm-as-judge|轨迹|回归|observability|可观测'),
 ('ai_infra','技术','AI Infra / MLOps / 推理','推理|vllm|tensorrt|gpu|cuda|k8s|kubernetes|量化|kv cache|mlops|部署|serving|模型服务'),
 ('python','技术','Python / AI 工程','python|fastapi|asyncio|pydantic'),
 ('java','技术','Java / 后端 / 分布式','java|spring|jvm|并发|线程池|mysql|redis|kafka|rocketmq|微服务|分布式|高并发'),
 ('architecture','技术','架构 / 系统设计','架构|系统设计|高可用|扩展性|稳定性|容灾|限流|熔断|领域建模|ddd|技术选型|重构'),
 ('payment_risk','技术','支付 / 风控 domain','支付|清结算|对账|资金|风控|规则引擎|决策引擎|反欺诈|合规'),
 ('data_platform','技术','数据平台','数据平台|数仓|实时计算|flink|spark|特征平台|数据治理'),
 ('devops','技术','工程效能 / DevOps','ci/cd|devops|sre|监控|告警|灰度|发布'),
 ('team_mgmt','管理','团队管理','团队管理|带团队|管理经验|团队负责人|团队建设|梯队|培养'),
 ('hiring_perf','管理','招聘 / 绩效','招聘|绩效|考核|晋升|淘汰|pip'),
 ('tech_leadership','管理','技术领导力 / 规划','技术规划|技术战略|roadmap|架构决策|技术选型|技术债|技术路线'),
 ('delivery','管理','项目交付 / 跨部门','项目管理|交付|进度|跨部门|协同|风险'),
 ('business','管理','业务结果 / 成本','业务结果|业务指标|roi|商业化|成本|预算|收益'),
 ('org_design','管理','组织设计 / 文化','组织设计|组织架构|文化建设|人才梯队|跨国|跨时区'),
 ('stability','管理','稳定性 / 应急','稳定性|故障|应急|容灾|on-call|演练|变更'),
 ('upward','管理','向上管理 / 影响力','汇报|向上管理|沟通|影响力|stakeholder'),
]
DIRS=['AI/Agent','管理/负责人','AI/算法/平台','架构/专家','Java/后端','支付/风控','其他']

def text_of(r):
    return ' '.join(str(r.get(k,'')) for k in ['responsibilities','requirements_full','nice_to_have']).lower()

def main():
    ap=argparse.ArgumentParser(description='Build the capability demand map (能力图谱.xlsx + reports/10-技术与管理能力图谱.md) from <data-dir>/processed/jobs_structured.csv')
    jp.add_argument(ap)
    data=jp.resolve(ap.parse_args().data_dir,produce=False)
    PROC,OUT=data/'processed',data/'reports'
    src=PROC/'jobs_structured.csv'
    if not src.exists():
        raise SystemExit(f'{src} not found: run organize_jobs.py (same --data-dir) first')
    OUT.mkdir(parents=True,exist_ok=True)
    s=pd.read_csv(src,dtype=str).fillna('')
    rel=s[s.is_relevant=='是'].copy()
    if not len(rel):
        raise SystemExit(f'no relevant jobs (is_relevant=是) in {src}')
    cand=rel[rel.senior_fit.astype(str).str.lower().isin(['true','是','1'])].copy()
    n_rel,n_cand=len(rel),len(cand)
    print('relevant',n_rel,'senior',n_cand)
    # counts
    rows=[]
    for cid,cat,name,kw in CAPS:
        pat=re.compile(kw,re.I)
        cnt=sum(1 for t in rel.responsibilities+' '+rel.requirements_full+' '+rel.nice_to_have if pat.search(t))
        cnt_c=sum(1 for t in cand.responsibilities+' '+cand.requirements_full+' '+cand.nice_to_have if pat.search(t))
        rows.append({'id':cid,'category':cat,'capability':name,
                     'market_count':cnt,'market_pct':round(cnt/n_rel*100,1),
                     'senior_count':cnt_c,'senior_pct':round(cnt_c/max(1,n_cand)*100,1)})
    cm=pd.DataFrame(rows)
    # direction x capability counts
    mat=[]
    for cid,cat,name,kw in CAPS:
        pat=re.compile(kw,re.I)
        row={'capability':name,'category':cat}
        for d in DIRS:
            g=rel[rel.direction==d]; n=sum(1 for t in g.responsibilities+' '+g.requirements_full+' '+g.nice_to_have if pat.search(t))
            row[d]=n
        mat.append(row)
    matrix=pd.DataFrame(mat)
    # job tags
    tags=[]
    for _,r in rel.iterrows():
        t=text_of(r); matched=[]
        for cid,cat,name,kw in CAPS:
            if re.search(kw,t,re.I): matched.append(name)
        tags.append({'job_id':r.job_id,'title':r.title,'direction':r.direction,'city':r.city,'senior_fit':r.senior_fit,'top_capabilities':'|'.join(matched[:8])})
    tagdf=pd.DataFrame(tags)
    # xlsx
    xlsx=PROC/'能力图谱.xlsx'
    with pd.ExcelWriter(xlsx,engine='openpyxl') as w:
        pd.DataFrame({
            '项目':['数据','岗位','口径','说明'],
            '值':[f'{n_rel} 条相关岗位 / {n_cand} 条 senior/经理/架构',f'技术+管理能力 {len(CAPS)} 项','频率=JD 文本命中岗位数/相关岗位数（正则关键词匹配）','关键词匹配较粗，用于判断方向与相对需求，不是岗位硬性门槛']
        }).to_excel(w,sheet_name='00-说明',index=False)
        cm.to_excel(w,sheet_name='01-能力图谱',index=False)
        matrix.to_excel(w,sheet_name='02-方向x能力',index=False)
        cm[cm.category=='管理'].to_excel(w,sheet_name='03-管理能力',index=False)
        cm[cm.category=='技术'].to_excel(w,sheet_name='04-技术能力',index=False)
        tagdf.to_excel(w,sheet_name='05-岗位能力标签',index=False)
    # style
    wb=load_workbook(xlsx)
    for ws in wb.worksheets:
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for c in ws[1]:
            c.fill=PatternFill('solid',fgColor='1F4E78'); c.font=Font(color='FFFFFF',bold=True)
        for col in ws.columns:
            width=max(12,min(50,max(len(str(c.value or '')) for c in col)+2))
            ws.column_dimensions[col[0].column_letter].width=width
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment=Alignment(vertical='top',wrap_text=True)
    wb.save(xlsx)
    # markdown
    mind=['```mermaid','mindmap','  root((技术+管理能力))']
    for cat in ['技术','管理']:
        mind.append('    '+cat)
        mind+=['      '+name for _,c,name,_ in CAPS if c==cat]
    mind.append('```')
    L=[f'# 技术 + 管理能力需求图谱（基于 {n_rel} 条相关 JD）','',
       f'> 数据：`processed/jobs_structured.csv`（{n_rel} 相关 / {n_cand} senior）。频率 = JD 职责/要求文本命中该能力的岗位数 / {n_rel}。',
       '> 关键词为正则匹配，粒度较粗，用于判断方向与相对需求。','',
       '## 0. 一页图谱','']+mind+['']
    for title,cat in [('## 1. 技术能力需求','技术'),('## 2. 管理能力需求','管理')]:
        L+=[title,'','| 能力 | 市场需求 | senior 岗需求 |','|---|---:|---:|']
        for _,r in cm[cm.category==cat].sort_values('market_count',ascending=False).iterrows():
            L.append(f"| {r.capability} | {r.market_count}（{r.market_pct}%） | {r.senior_count}（{r.senior_pct}%） |")
        L.append('')
    L+=['## 3. 方向 × 能力矩阵（命中岗位数）','']
    # markdown matrix compact
    cols=['capability']+DIRS[:-1]
    L.append('| 能力 | '+' | '.join(cols[1:])+' |')
    L.append('|---|'+'|'.join(['---:']*(len(cols)-1))+'|')
    for _,r in matrix[cols].iterrows():
        L.append('| '+r.capability+' | '+' | '.join(str(int(r[c])) for c in cols[1:])+' |')
    L+=['','## 4. 文件','',
        '- `processed/能力图谱.xlsx`',
        '- `processed/jobs_structured.csv`',
        '- `reports/10-技术与管理能力图谱.md`']
    (OUT/'10-技术与管理能力图谱.md').write_text('\n'.join(L)+'\n',encoding='utf-8')
    print('wrote',OUT/'10-技术与管理能力图谱.md',xlsx)

if __name__=='__main__': main()
