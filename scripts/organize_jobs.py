#!/usr/bin/env python3
# Organize all collected jobs into a multi-sheet workbook + structured CSV (all under <data-dir>/processed).
import argparse
import re
from collections import Counter
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import jobintel_paths as jp

SKILLS=[
 ('Agent','Agent'),('智能体','Agent'),('Multi-Agent','Multi-Agent'),('MCP','MCP'),('LangGraph','LangGraph'),
 ('RAG','RAG'),('LLM','LLM'),('大模型','大模型'),('Prompt','Prompt'),('上下文','Context Engineering'),
 ('Function Calling','Function Calling'),('工具调用','Tool Calling'),('Harness','Harness'),('Runtime','Runtime'),
 ('评测','Evaluation'),('Evaluation','Evaluation'),('推理','推理'),('微调','微调'),('RLHF','RLHF'),
 ('Java','Java'),('Spring','Spring'),('JVM','JVM'),('并发','并发'),('分布式','分布式'),('微服务','微服务'),
 ('高并发','高并发'),('MySQL','MySQL'),('Redis','Redis'),('Kafka','Kafka'),('RocketMQ','RocketMQ'),
 ('Elasticsearch','Elasticsearch'),('团队管理','团队管理'),('架构','架构'),('支付','支付'),('清结算','清结算'),
 ('风控','风控'),('合规','合规'),('数据平台','数据平台'),('MLOps','MLOps'),('K8s','K8s'),('云原生','云原生'),
 ('Python','Python'),('Go','Go'),('C++','C++'),('算法','算法'),('多模态','多模态'),('搜索','搜索'),('推荐','推荐'),
]
HEAD_RESP=re.compile(r'(岗位职责|工作职责|职责描述|职位职责|工作内容|你将负责|岗位描述|职位描述|主要职责)')
HEAD_REQ=re.compile(r'(任职要求|岗位要求|任职资格|职位要求|能力要求|我们希望你|任职条件|岗位基本要求)')
HEAD_BONUS=re.compile(r'(加分项|优先条件|以下条件优先|具备以下条件优先|加分)')
HEAD_END=re.compile(r'(公司介绍|关于我们|团队介绍|福利待遇|工作地址|联系方式|竞争力分析|相似职位|热门职位|工商信息)')


def clean(s):
    s=str(s or '')
    s=re.sub(r'boss\s*直聘|BOSS\s*直聘','',s,flags=re.I)
    s=re.sub(r'直聘|boss','',s,flags=re.I)
    s=re.sub(r'\s+',' ',s)
    s=re.sub(r'^[\s:：,，;；、|]+','',s)
    return s.strip()


def section(text, head, next_heads):
    m=head.search(text)
    if not m: return ''
    start=m.end()
    tail=text[start:]
    ends=[len(tail)]
    for h in next_heads:
        mm=h.search(tail)
        if mm: ends.append(mm.start())
    return clean(tail[:min(ends)])


def extract_sections(jd):
    t=clean(jd)
    resp=section(t,HEAD_RESP,[HEAD_REQ,HEAD_BONUS,HEAD_END]) or ''
    req=section(t,HEAD_REQ,[HEAD_BONUS,HEAD_END]) or ''
    bonus=section(t,HEAD_BONUS,[HEAD_END]) or ''
    if not resp and req:
        m=HEAD_REQ.search(t); resp=clean(t[:m.start()]) if m else ''
    if not resp and not req:
        resp=t
    # if headings repeated/empty, fallback whole
    return resp[:2500], req[:2500], bonus[:1200]


def extract_skills(text):
    low=str(text).lower(); out=[]
    for key,label in SKILLS:
        if key.lower() in low and label not in out:
            out.append(label)
    return '|'.join(out[:25])


def hard_requirements(req):
    if not req: return ''
    parts=re.split(r'(?<=[。；;])\s*|\n',req)
    keep=[]
    for p in parts:
        p=clean(p)
        if not p: continue
        if re.search(r'必须|至少|精通|熟悉|掌握|年以上|本科|硕士|博士|有.{0,12}经验|要求',p):
            keep.append(p)
    return '；'.join(keep[:10])


def management_scope(text):
    if not text: return ''
    m=re.search(r'(?:管理过?|带|团队规模|管理)\s*(\d+)\s*(?:-\s*(\d+))?\s*人',text)
    if m:
        return f"{m.group(1)}-{m.group(2)}人" if m.group(2) else f"{m.group(1)}人"
    if '团队管理' in text or '管理经验' in text or '带团队' in text:
        return '有团队管理要求（人数未写）'
    return ''


def years_required(text):
    if not text: return ''
    years=[int(x) for x in re.findall(r'(\d+)\s*年(?:以上|\+)?',text)]
    return str(max(years))+'年+' if years else ''


def education(text):
    for x in ['博士','硕士','本科','大专','学历不限']:
        if x in str(text): return x
    return ''


def sub_direction(title, fam):
    t=str(title); low=t.lower()
    if fam=='AI/Agent':
        if any(k in low for k in ['平台','infra','infrastructure','mlops','runtime','harness']): return 'Agent平台/Infra'
        if '架构' in t: return 'Agent架构'
        if any(k in low for k in ['rag','检索','知识库']): return 'RAG/知识库'
        if any(k in low for k in ['评测','evaluation']): return 'Agent评测'
        return 'Agent开发/应用'
    if fam=='Java/后端':
        if '架构' in t: return 'Java/后端架构'
        return 'Java/后端开发'
    if fam=='管理/负责人':
        if '总监' in t or 'cto' in low: return '技术总监/CTO'
        if '负责人' in t: return '技术负责人'
        return '技术经理/研发经理'
    if fam=='架构/专家':
        return '架构师/技术专家'
    if fam=='AI/算法/平台':
        if any(k in low for k in ['平台','mlops','推理','infra']): return 'AI平台/Infra'
        return 'AI应用/算法'
    if fam=='支付/风控':
        return '支付/风控技术'
    return '其他'


def level_guess(title, senior_fit):
    t=str(title)
    if any(k in t for k in ['CTO','首席']): return 'CTO/首席'
    if '总监' in t: return '总监'
    if '负责人' in t: return '负责人'
    if '经理' in t: return '经理'
    if '架构' in t: return '架构师'
    if '专家' in t: return '专家'
    if any(k in t for k in ['资深','高级','Staff','Principal']): return '高级/资深'
    return '高级/资深' if str(senior_fit).lower()=='true' else '中级/其他'


def style_sheet(ws, wide_cols=(), freeze='A2'):
    ws.freeze_panes=freeze
    ws.auto_filter.ref=ws.dimensions
    header_fill=PatternFill('solid',fgColor='1F4E78')
    for c in ws[1]:
        c.fill=header_fill; c.font=Font(color='FFFFFF',bold=True); c.alignment=Alignment(vertical='center')
    for i,col in enumerate(ws.iter_cols(min_row=1,max_row=1),1):
        name=col[0].value or ''
        width=40 if name in wide_cols else min(max(12,len(str(name))+2),24)
        ws.column_dimensions[get_column_letter(i)].width=width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment=Alignment(vertical='top',wrap_text=(c.column_letter in [get_column_letter(i) for i in range(1,ws.max_column+1)]))


def main():
    ap=argparse.ArgumentParser(description='Organize <data-dir>/processed/jobs_master_merged.csv + jobs_filtered.csv into jobs_structured.csv and 岗位信息整理.xlsx')
    jp.add_argument(ap)
    PROC=jp.resolve(ap.parse_args().data_dir,produce=False)/'processed'
    merged,filtered=PROC/'jobs_master_merged.csv',PROC/'jobs_filtered.csv'
    if not merged.exists():
        raise SystemExit(f'{merged} not found: run analyze.py (same --data-dir) first')
    if not filtered.exists():
        raise SystemExit(f'{filtered} not found: jobs_filtered.csv / jobs_removed.csv come from a filter step that is not part of this package (see WORKFLOW.md §3)')
    all_df=pd.read_csv(merged,dtype=str).fillna('')
    filt=pd.read_csv(filtered,dtype=str).fillna('')
    rem_path=PROC/'jobs_removed.csv'
    rem=pd.read_csv(rem_path,dtype=str).fillna('') if rem_path.exists() else pd.DataFrame(columns=['job_id','filter_reason'])
    keep_ids=set(filt.job_id)
    senior_ids=set(filt[filt.senior_fit.astype(str).str.lower().isin(['true','是','1'])].job_id)
    reason_map=dict(zip(rem.job_id,rem.filter_reason))
    rows=[]
    for _,r in all_df.iterrows():
        jd=r.get('jd_text','')
        resp,req,bonus=extract_sections(jd)
        text=' '.join([resp,req,bonus])
        fam=r.get('role_family','') or '其他'
        senior= r['job_id'] in senior_ids
        rows.append({
            'job_id':r['job_id'],'source_platform':r.get('source_platform',''),
            'company':r.get('company',''),'title':r.get('title',''),
            'direction':fam,'sub_direction':sub_direction(r.get('title',''),fam),
            'level_guess':level_guess(r.get('title',''),senior),
            'is_relevant':'是' if r['job_id'] in keep_ids else '否',
            'senior_fit':'是' if senior else '否',
            'filter_reason':reason_map.get(r['job_id'],''),
            'city':r.get('city',''),'salary_raw':r.get('salary_raw',''),
            'salary_min_cny':r.get('salary_min_cny',''),'salary_max_cny':r.get('salary_max_cny',''),
            'experience':r.get('experience',''),'education':r.get('education',''),
            'years_required':years_required(req+resp),
            'management_scope':management_scope(req+resp),
            'must_have_skills':extract_skills(req+resp),
            'responsibilities':resp,'hard_requirements':hard_requirements(req),
            'requirements_full':req,'nice_to_have':bonus,
            'jd_quality':('完整' if len(resp)>=30 and len(req)>=30 else '部分' if (len(resp)>=30 or len(req)>=30) else '短/缺失'),
            'source_url':r.get('source_url',''),'verification':r.get('verification',''),
            'collected_at':r.get('collected_at',''),
        })
    s=pd.DataFrame(rows)
    s.to_csv(PROC/'jobs_structured.csv',index=False,encoding='utf-8-sig')
    rel=s[s.is_relevant=='是'].copy(); cand=s[(s.is_relevant=='是')&(s.senior_fit=='是')].copy()
    print('all',len(s),'relevant',len(rel),'candidate',len(cand))
    print('jd quality'); print(s.jd_quality.value_counts().to_string())
    print('sections with req', (s.requirements_full.str.len()>=20).sum(), 'resp',(s.responsibilities.str.len()>=20).sum())

    cols=['company','title','direction','sub_direction','level_guess','city','salary_raw','experience','education','years_required','management_scope','must_have_skills','responsibilities','hard_requirements','requirements_full','nice_to_have','source_platform','source_url']
    xlsx=PROC/'岗位信息整理.xlsx'
    with pd.ExcelWriter(xlsx,engine='openpyxl') as w:
        # README
        info=pd.DataFrame({
            '项目':['数据源','总岗位','相关岗位','投递候选(senior/经理/架构)','已删除(二次复核后)','JD完整','说明'],
            '值':[' + '.join(sorted({x for x in s.source_platform if x})),len(s),len(rel),len(cand),len(rem),int((s.jd_quality=='完整').sum()),'按 jobs_filtered.csv / jobs_removed.csv 标记相关岗位；按方向/级别/城市整理，职责/要求从 JD 自动拆分']
        })
        info.to_excel(w,sheet_name='00-说明',index=False)
        s.to_excel(w,sheet_name='01-全部岗位',index=False)
        rel[cols].to_excel(w,sheet_name=f'02-相关岗位-{len(rel)}',index=False)
        cand[cols].to_excel(w,sheet_name=f'03-投递候选-{len(cand)}',index=False)
        for fam,sheet in [('AI/Agent','04-AI-Agent'),('Java/后端','05-Java后端'),('管理/负责人','06-管理负责人'),('架构/专家','07-架构专家'),('AI/算法/平台','08-AI算法平台'),('支付/风控','09-支付风控')]:
            d=rel[rel.direction==fam][cols]
            if len(d): d.to_excel(w,sheet_name=sheet,index=False)
        # map
        mp=pd.pivot_table(rel,index='direction',columns='city',values='job_id',aggfunc='count',fill_value=0)
        mp['合计']=mp.sum(axis=1)
        mp.to_excel(w,sheet_name='10-岗位地图')
        # skills by direction
        sk_rows=[]
        for fam,g in rel.groupby('direction'):
            c=Counter()
            for x in g.must_have_skills.fillna(''):
                for k in str(x).split('|'):
                    if k: c[k]+=1
            for k,n in c.most_common(12): sk_rows.append({'direction':fam,'skill':k,'count':n,'pct':round(n/max(1,len(g))*100,1)})
        pd.DataFrame(sk_rows).to_excel(w,sheet_name='11-技能矩阵',index=False)
        # removed
        if len(rem):
            rem[['company','title','role_family','city','salary_raw','filter_reason','source_platform','source_url']].to_excel(w,sheet_name='12-剔除岗位',index=False)
    # style
    wb=load_workbook(xlsx)
    for ws in wb.worksheets:
        style_sheet(ws,wide_cols={'responsibilities','hard_requirements','requirements_full','nice_to_have','must_have_skills'})
    wb.save(xlsx)
    print('wrote',xlsx)

if __name__=='__main__': main()
